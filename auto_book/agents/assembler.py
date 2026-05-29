"""Assemblers for Markdown and DOCX book exports."""

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from auto_book.models.book_bible import BookBible
from auto_book.models.chapter import ChapterDraft
from auto_book.models.export import ExportResult
from auto_book.models.image import ImageAsset
from auto_book.utils.logger import get_logger

IMAGE_ANCHOR_PATTERN = re.compile(r"^\[IMAGE_ANCHOR:\s*([A-Z0-9_]+)\]\s*$")

_DARK_BLUE = RGBColor(31, 56, 100)
_GREY = RGBColor(128, 128, 128)


def assemble_markdown(
    book_bible: BookBible,
    chapters: list[ChapterDraft],
    output_dir: str,
    image_assets: list[ImageAsset] | None = None,
) -> ExportResult:
    """Assemble accepted chapters into output/book.md with image references."""

    logger = get_logger()
    output_path = Path(output_dir) / "book.md"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    ordered = sorted(chapters, key=lambda chapter: chapter.chapter_number)
    images = image_assets or []
    parts: list[str] = []

    parts.append(f"# {book_bible.working_title}\n")
    if book_bible.subtitle:
        parts.append(f"### {book_bible.subtitle}\n")
    parts.append(f"*{book_bible.book_promise}*\n")
    parts.append("---\n")

    parts.append("## Table of Contents\n")
    for chapter in ordered:
        parts.append(
            f"- [Chapter {chapter.chapter_number}: {chapter.title}]"
            f"(#{_anchor(chapter.title)})"
        )
    parts.append("\n---\n")

    for chapter in ordered:
        parts.append(f"\n## Chapter {chapter.chapter_number}: {chapter.title}\n")
        parts.append(
            _replace_markdown_anchors(
                _strip_duplicate_title(chapter),
                images,
                output_path.parent,
            )
        )
        if chapter.key_takeaways:
            parts.append("\n### Key Takeaways\n")
            parts.extend(f"- {takeaway}" for takeaway in chapter.key_takeaways)
        parts.append("\n---\n")

    content = "\n".join(parts).strip() + "\n"
    output_path.write_text(content, encoding="utf-8")
    word_count = len(content.split())
    logger.info(
        "Markdown assembled at %s (%s chapters, %s words)",
        output_path,
        len(ordered),
        word_count,
    )
    return ExportResult(
        markdown_path=str(output_path),
        total_chapters=len(ordered),
        total_word_count=word_count,
        images_embedded=sum(1 for asset in images if not asset.is_placeholder),
        success=True,
    )


def assemble_docx(
    book_bible: BookBible,
    chapters: list[ChapterDraft],
    output_dir: str,
    image_assets: list[ImageAsset] | None = None,
    cover_asset: ImageAsset | None = None,
) -> str:
    """Assemble accepted chapters and available images into output/book.docx."""

    logger = get_logger()
    output_path = Path(output_dir) / "book.docx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    ordered = sorted(chapters, key=lambda chapter: chapter.chapter_number)
    images = image_assets or []

    template = Path("resources/template.docx")
    has_template = False
    doc = Document()
    if template.exists():
        try:
            t = Document(str(template))
            from collections import Counter
            dupes = [n for n, c in Counter(s.name for s in t.styles).items() if c > 1]
            if dupes:
                get_logger().warning(
                    "Template has duplicate styles %s — fix in Word and re-save. Using blank document.",
                    dupes,
                )
            else:
                doc = t
                has_template = True
        except Exception as exc:
            get_logger().warning("Template load failed (%s); using blank document.", exc)

    _add_title_page(doc, book_bible, has_template, cover_asset)
    doc.add_page_break()
    _add_toc(doc, chapters=ordered, has_template=has_template)
    _add_header_footer(doc, book_bible)
    for chapter in ordered:
        doc.add_page_break()
        _add_chapter(doc, chapter, images, has_template)

    doc.save(str(output_path))
    logger.info("DOCX assembled at %s", output_path)
    return str(output_path)


def _add_title_page(
    doc: Document,
    bible: BookBible,
    has_template: bool = False,
    cover_asset: ImageAsset | None = None,
) -> None:
    # Cover image — full page width if available
    if cover_asset and cover_asset.file_path and Path(cover_asset.file_path).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(cover_asset.file_path, width=Inches(6.5))
        doc.add_paragraph("")
    else:
        for _ in range(5):
            doc.add_paragraph("")

    # Title
    if has_template:
        title = doc.add_paragraph(bible.working_title, style="Title")
    else:
        title = doc.add_heading(bible.working_title, level=0)
        for run in title.runs:
            run.font.size = Pt(28)
            run.font.bold = True
            run.font.color.rgb = _DARK_BLUE
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Subtitle
    if bible.subtitle:
        if has_template:
            subtitle = doc.add_paragraph(bible.subtitle, style="Subtitle")
        else:
            subtitle = doc.add_paragraph(bible.subtitle)
            for run in subtitle.runs:
                run.font.size = Pt(16)
                run.font.italic = True
                run.font.color.rgb = RGBColor(100, 100, 100)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")
    tagline = doc.add_paragraph(bible.book_promise)
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in tagline.runs:
        run.font.italic = True
        run.font.size = Pt(12)
        if not has_template:
            run.font.color.rgb = RGBColor(80, 80, 80)


def _add_toc(doc: Document, chapters: list | None = None, has_template: bool = False) -> None:
    """Add a static Table of Contents with chapter titles and section subheadings."""
    if has_template:
        try:
            doc.add_paragraph("Table of Contents", style="TOC Heading")
        except Exception:
            doc.add_heading("Table of Contents", level=1)
    else:
        h = doc.add_heading("Table of Contents", level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in h.runs:
            run.font.size = Pt(20)
            run.font.color.rgb = _DARK_BLUE

    if not chapters:
        return

    for chapter in chapters:
        # Chapter entry (toc 1)
        try:
            p = doc.add_paragraph(style="toc 1" if has_template else "Normal")
        except Exception:
            p = doc.add_paragraph()
        run = p.add_run(f"Chapter {chapter.chapter_number}: {chapter.title}")
        run.font.bold = True
        if not has_template:
            run.font.size = Pt(11)
            run.font.color.rgb = _DARK_BLUE
        dots = p.add_run("  " + "." * 35)
        dots.font.size = Pt(9)
        dots.font.color.rgb = _GREY
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(6)

        # Section subheadings (toc 2) — extract ## lines from chapter body
        for line in chapter.body.splitlines():
            stripped = line.strip()
            if stripped.startswith("## "):
                section_title = stripped[3:].strip()
                try:
                    sp = doc.add_paragraph(style="toc 2" if has_template else "Normal")
                except Exception:
                    sp = doc.add_paragraph()
                sr = sp.add_run(f"    {section_title}")
                if not has_template:
                    sr.font.size = Pt(10)
                    sr.font.color.rgb = _GREY
                sp.paragraph_format.space_after = Pt(1)
                sp.paragraph_format.left_indent = Pt(18)


def _add_header_footer(doc: Document, bible: BookBible) -> None:
    """Add book title in header and page number in footer across all sections."""
    for section in doc.sections:
        section.different_first_page_header_footer = True

        header = section.header
        if not header.paragraphs:
            header.add_paragraph()
        hp = header.paragraphs[0]
        hp.clear()
        run = hp.add_run(bible.working_title)
        run.font.size = Pt(9)
        run.font.color.rgb = _GREY
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        footer = section.footer
        if not footer.paragraphs:
            footer.add_paragraph()
        fp = footer.paragraphs[0]
        fp.clear()
        _add_page_number_field(fp)


def _add_page_number_field(paragraph) -> None:
    """Inject a Word PAGE field into a paragraph."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for tag, field_text in [("begin", None), (None, " PAGE "), ("end", None)]:
        run = paragraph.add_run()
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = field_text
            run._r.append(el)
        run.font.size = Pt(9)
        run.font.color.rgb = _GREY


def _add_chapter(
    doc: Document,
    chapter: ChapterDraft,
    image_assets: list[ImageAsset],
    has_template: bool = False,
) -> None:
    heading = doc.add_heading(f"Chapter {chapter.chapter_number}: {chapter.title}", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        run.font.size = Pt(22)
        run.font.bold = True
        if not has_template:
            run.font.color.rgb = _DARK_BLUE
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(16)

    _add_markdown_body(doc, _strip_duplicate_title(chapter), image_assets, has_template)

    if chapter.key_takeaways:
        doc.add_heading("Key Takeaways", level=3)
        for takeaway in chapter.key_takeaways:
            p = doc.add_paragraph(takeaway, style="List Bullet")
            if not has_template:
                p.paragraph_format.space_after = Pt(3)


def _add_markdown_body(
    doc: Document,
    markdown: str,
    image_assets: list[ImageAsset],
    has_template: bool = False,
) -> None:
    for raw_line in markdown.splitlines():
        line = raw_line.strip()

        # Blank line = paragraph break spacer
        if not line:
            if not has_template:
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(4)
            continue

        anchor_match = IMAGE_ANCHOR_PATTERN.fullmatch(line)
        if anchor_match:
            _add_docx_anchor_image(doc, anchor_match.group(1), image_assets)
        elif line.startswith("### "):
            h = doc.add_heading(line[4:], level=3)
            if not has_template:
                h.paragraph_format.space_before = Pt(10)
        elif line.startswith("## "):
            h = doc.add_heading(line[3:], level=2)
            if not has_template:
                for run in h.runs:
                    run.font.size = Pt(14)
                    run.font.bold = True
                h.paragraph_format.space_before = Pt(14)
                h.paragraph_format.space_after = Pt(4)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(line[2:], style="List Bullet")
            if not has_template:
                p.paragraph_format.space_after = Pt(3)
        elif re.match(r"^\d+[\.)]\s+", line):
            p = doc.add_paragraph(re.sub(r"^\d+[\.)]\s+", "", line), style="List Number")
            if not has_template:
                p.paragraph_format.space_after = Pt(3)
        else:
            paragraph = doc.add_paragraph()
            _add_formatted_text(paragraph, line)
            if not has_template:
                paragraph.paragraph_format.first_line_indent = Pt(18)
                paragraph.paragraph_format.space_after = Pt(10)
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    run.font.size = Pt(11)


def _add_formatted_text(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def _replace_markdown_anchors(
    markdown: str,
    image_assets: list[ImageAsset],
    markdown_dir: Path,
) -> str:
    lines: list[str] = []
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        match = IMAGE_ANCHOR_PATTERN.fullmatch(line)
        if not match:
            lines.append(raw_line)
            continue
        lines.append(_markdown_for_anchor(match.group(1), image_assets, markdown_dir))
    return "\n".join(lines).strip()


def _markdown_for_anchor(
    anchor_id: str,
    image_assets: list[ImageAsset],
    markdown_dir: Path,
) -> str:
    image = _image_for_anchor(anchor_id, image_assets)
    if image is None or image.is_placeholder or not image.file_path:
        label = (
            (image.alt_text if image else "")
            or (image.prompt_used if image else "")
            or anchor_id.replace("_", " ").title()
        )
        return f"*[Image placeholder: {label}]*"

    image_path = Path(image.file_path)
    try:
        display_path = image_path.relative_to(markdown_dir)
    except ValueError:
        display_path = image_path
    alt = image.alt_text or anchor_id.replace("_", " ").title()
    return f"![{alt}]({display_path.as_posix()})"


def _add_docx_anchor_image(
    doc: Document,
    anchor_id: str,
    image_assets: list[ImageAsset],
) -> None:
    image = _image_for_anchor(anchor_id, image_assets)
    if image is None or image.is_placeholder or not image.file_path:
        label = (
            (image.alt_text if image else "")
            or (image.prompt_used if image else "")
            or anchor_id.replace("_", " ").title()
        )
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"[Image placeholder: {label}]")
        run.italic = True
        return

    image_path = Path(image.file_path)
    if not image_path.exists():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"[Image missing: {image.alt_text or anchor_id}]")
        run.italic = True
        return

    # Full-width image centered
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(5.5))

    # Caption
    if image.alt_text:
        caption = doc.add_paragraph(image.alt_text)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(10)
        for run in caption.runs:
            run.font.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = _GREY


def _image_for_anchor(
    anchor_id: str,
    image_assets: list[ImageAsset],
) -> ImageAsset | None:
    for image in image_assets:
        if image.anchor_id == anchor_id:
            return image
    return None


def _anchor(title: str) -> str:
    anchor = title.strip().lower()
    anchor = re.sub(r"[^a-z0-9\s-]", "", anchor)
    anchor = re.sub(r"\s+", "-", anchor)
    return anchor


def _strip_duplicate_title(chapter: ChapterDraft) -> str:
    lines = chapter.body.splitlines()
    if lines and lines[0].strip().startswith("#") and chapter.title.lower() in lines[0].lower():
        return "\n".join(lines[1:]).strip()
    return chapter.body.strip()
